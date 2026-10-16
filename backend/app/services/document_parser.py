import io
import re
from typing import Any, Dict, List, Optional


class DocumentExtractionError(Exception):
    """Raised when a document cannot be opened or read (corrupt, encrypted, unsupported)."""


class ParsedQuestion:
    def __init__(
        self,
        question_number: str,
        raw_text: str,
        clean_text: str,
        section_name: Optional[str] = None,
        parent_question_number: Optional[str] = None,
        marks: int = 0
    ):
        self.question_number = question_number
        self.raw_text = raw_text
        self.clean_text = clean_text
        self.section_name = section_name
        self.parent_question_number = parent_question_number
        self.marks = marks

    def to_dict(self) -> Dict[str, Any]:
        return {
            "question_number": self.question_number,
            "raw_text": self.raw_text,
            "clean_text": self.clean_text,
            "section_name": self.section_name,
            "parent_question_number": self.parent_question_number,
            "marks": self.marks,
        }


SECTION_REGEX = re.compile(r"^(?:SECTION|PART)\s+([A-Z0-9]+)\b", re.IGNORECASE)
# "Q1.", "Question 4:", "1.", "2)" -- but not decimals like "1.5 kg" (negative lookahead for a digit).
QUESTION_REGEX = re.compile(r"^(?:Q(?:uestion)?\s*)?(\d{1,3})\s*[\.\:\)](?!\d)\s*(.*)", re.IGNORECASE)
# Lower-case sub-parts only: "(a)", "a)", "a.", "(ii)". Upper-case "A." / "(A)" are MCQ options and stay in the stem.
SUBQUESTION_REGEX = re.compile(r"^(?:\(([a-h]|[ivx]{1,4})\)|([a-h])[\.\)])\s+(.*)")
MARKS_REGEX = re.compile(r"[\(\[]\s*(\d{1,3})\s*(?:marks?|mks?|pts?|points?)\s*[\)\]]", re.IGNORECASE)
TRAILING_BRACKET_MARKS_REGEX = re.compile(r"\[\s*(\d{1,3})\s*\]\s*$")


def _find_marks(line: str) -> Optional[int]:
    m = MARKS_REGEX.search(line) or TRAILING_BRACKET_MARKS_REGEX.search(line)
    return int(m.group(1)) if m else None


class DocumentParserService:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> tuple[str, bool]:
        """Extract text from a PDF with PyMuPDF. Returns (text, needs_ocr)."""
        import pymupdf as fitz

        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as exc:
            raise DocumentExtractionError("The PDF file is corrupt or unreadable.") from exc
        with doc:
            if doc.needs_pass:
                raise DocumentExtractionError("File is encrypted. Please upload an unprotected PDF/DOCX.")
            # sort=True orders text blocks top-left to bottom-right, which helps multi-column layouts.
            text = "\n".join(page.get_text("text", sort=True) for page in doc).strip()
        return text, len(text) < 50

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """Extract paragraph and table text, synthesising Word auto-numbering ("1.", "(a)")."""
        import docx
        from docx.oxml.ns import qn

        try:
            document = docx.Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise DocumentExtractionError("The DOCX file is corrupt or unreadable.") from exc

        lines: List[str] = []
        top_counter = 0
        sub_counter = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            num_pr = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
            style_name = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if num_pr is not None or "list number" in style_name:
                level = 0
                if num_pr is not None:
                    ilvl = num_pr.find(qn("w:ilvl"))
                    if ilvl is not None:
                        level = int(ilvl.get(qn("w:val"), "0"))
                if level == 0 and not QUESTION_REGEX.match(text):
                    top_counter += 1
                    sub_counter = 0
                    text = f"{top_counter}. {text}"
                elif level > 0 and not SUBQUESTION_REGEX.match(text):
                    sub_counter += 1
                    text = f"({chr(ord('a') + (sub_counter - 1) % 26)}) {text}"
            lines.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" ".join(dict.fromkeys(cells)))
        return "\n".join(lines)

    @classmethod
    def parse_questions_from_text(cls, text: str) -> List[ParsedQuestion]:
        """Regex and heuristic layout parser for question segmentation."""
        questions: List[ParsedQuestion] = []

        current_section = "General"
        current_q_num: Optional[str] = None
        current_q_lines: List[str] = []
        current_marks = 0
        # Index into `questions` of the sub-question currently receiving continuation lines.
        current_sub_index: Optional[int] = None
        parent_insert_index = 0

        def flush_parent() -> None:
            nonlocal current_q_lines, current_marks
            if current_q_num is not None and current_q_lines:
                full_q_text = " ".join(current_q_lines).strip()
                # Insert the parent before its sub-questions to preserve reading order.
                questions.insert(
                    parent_insert_index,
                    ParsedQuestion(
                        question_number=current_q_num,
                        raw_text=full_q_text,
                        clean_text=cls._clean_text(full_q_text),
                        section_name=current_section,
                        # Marks can be split across wrapped lines, e.g. "(10" / "marks)".
                        marks=current_marks or _find_marks(full_q_text) or 0,
                    ),
                )
            current_q_lines = []
            current_marks = 0

        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue

            sec_match = SECTION_REGEX.match(stripped)
            if sec_match:
                flush_parent()
                current_q_num = None
                current_sub_index = None
                current_section = f"Section {sec_match.group(1).upper()}"
                continue

            q_match = QUESTION_REGEX.match(stripped)
            if q_match:
                flush_parent()
                current_q_num = q_match.group(1)
                current_sub_index = None
                parent_insert_index = len(questions)
                rest_of_line = q_match.group(2)
                if rest_of_line:
                    current_q_lines.append(rest_of_line)
                current_marks = _find_marks(stripped) or 0
                continue

            sub_match = SUBQUESTION_REGEX.match(stripped)
            if sub_match and current_q_num is not None:
                sub_label = sub_match.group(1) or sub_match.group(2)
                sub_text = sub_match.group(3)
                questions.append(
                    ParsedQuestion(
                        question_number=f"{current_q_num}({sub_label})",
                        raw_text=sub_text,
                        clean_text=cls._clean_text(sub_text),
                        section_name=current_section,
                        parent_question_number=current_q_num,
                        marks=_find_marks(stripped) or 0,
                    )
                )
                current_sub_index = len(questions) - 1
                continue

            if current_q_num is None:
                continue  # header/instructions before the first question

            if current_sub_index is not None:
                sub = questions[current_sub_index]
                sub.raw_text = f"{sub.raw_text} {stripped}"
                sub.clean_text = cls._clean_text(sub.raw_text)
                if not sub.marks:
                    sub.marks = _find_marks(sub.raw_text) or 0
            else:
                current_q_lines.append(stripped)
                if current_marks == 0:
                    current_marks = _find_marks(stripped) or 0

        flush_parent()

        # A sub-question whose parent had no stem text keeps its parent number; the pipeline
        # treats a missing parent as a top-level question.
        return [q for q in questions if q.clean_text]

    @staticmethod
    def _clean_text(text: str) -> str:
        cleaned = MARKS_REGEX.sub("", text)
        cleaned = TRAILING_BRACKET_MARKS_REGEX.sub("", cleaned)
        cleaned = re.sub(r"\s+", " ", cleaned).strip()
        return cleaned
